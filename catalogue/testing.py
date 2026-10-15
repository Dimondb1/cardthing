"""Small helpers for building test data."""

from datetime import timedelta
from decimal import Decimal

from django.utils import timezone

from .models import Game, Listing, Product, ProductSet, Retailer


def make_game(name="Pokémon", slug="pokemon", **kwargs):
    kwargs.setdefault("search_aliases", "ptcg")
    return Game.objects.create(name=name, slug=slug, **kwargs)


def make_set(game, name="Prismatic Evolutions", slug="prismatic-evolutions", **kwargs):
    kwargs.setdefault("code", "PRE")
    return ProductSet.objects.create(game=game, name=name, slug=slug, **kwargs)


def make_product(product_set, name="Prismatic Evolutions Elite Trainer Box", **kwargs):
    kwargs.setdefault("product_type", Product.Type.ELITE_TRAINER_BOX)
    return Product.objects.create(
        game=product_set.game, product_set=product_set, name=name, **kwargs
    )


def make_retailer(name="Harbour Games", **kwargs):
    slug = kwargs.pop("slug", name.lower().replace(" ", "-").replace(".", ""))
    kwargs.setdefault("website", f"https://www.{slug}.example/")
    return Retailer.objects.create(name=name, slug=slug, **kwargs)


def make_listing(product, retailer, price="50.00", delivery="0.00", hours_ago=1, url=None, **kwargs):
    kwargs.setdefault("availability", Listing.Availability.IN_STOCK)
    return Listing.objects.create(
        product=product,
        retailer=retailer,
        url=url or f"{retailer.website}p/{product.slug}",
        price=Decimal(price),
        delivery_cost=Decimal(delivery),
        last_checked=timezone.now() - timedelta(hours=hours_ago),
        **kwargs,
    )


def inside_the_cache_window(test):
    """
    Act as if this process cleared the list caches a moment ago, for the rest of the test.

    An unforced clear (a save) is then skipped for 30 seconds, so a clear the test sees can only be
    a forced one.
    """
    import time
    from unittest import mock

    from django.test import override_settings

    overrides = override_settings(RIPRAPTOR_CACHE_CLEAR_SECONDS=30)
    overrides.enable()
    test.addCleanup(overrides.disable)
    patcher = mock.patch("catalogue.signals._last_clear", time.monotonic())
    patcher.start()
    test.addCleanup(patcher.stop)


class Pasted:
    """A value pasted into a cell rather than typed, kept as it is: a formula, a number or a date."""

    def __init__(self, value):
        self.value = value


def excel_typed(text, number_format):
    """What Excel keeps when ``text`` is typed into a cell with this number format, on UK settings. A text
    cell keeps exactly what was typed, an = included. Any other cell makes digits a number (losing leading
    zeros), TRUE or FALSE a true or false, a day-first date a date, and = a formula with no value."""
    import re
    from datetime import datetime

    if number_format == "@":
        return text
    plain = text.strip()
    if plain.startswith("="):
        return Pasted(plain)
    if plain.upper() in ("TRUE", "FALSE"):
        return plain.upper() == "TRUE"
    if re.fullmatch(r"[-+]?\d+", plain):
        number = int(plain)
        # Excel keeps 15 significant digits.
        return number if len(plain.lstrip("+-0")) <= 15 else float(number)
    if re.fullmatch(r"[-+]?(?:\d+\.\d*|\.\d+|\d+)(?:[eE][-+]?\d+)?", plain):
        return float(plain)
    for pattern in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(plain, pattern)
        except ValueError:
            pass
    return text


def excel_resave(data, edits=(), *, drop_custom_props=False, drop_validations=False, sort=None, delete_rows=(),
                 insert_rows=(), reorder_columns=False, extra_column=False):
    """An exported workbook as it comes back after the owner opens it in Excel, changes it and saves it.

    ``edits`` are (sheet, row key, heading, value): the row is found by its id (or link, on Left out), and a
    str value is typed as Excel stores typing (excel_typed) while a Pasted value goes in as it is.
    ``insert_rows`` are (sheet, {heading: value}) rows added at the bottom, typed the same way, each cell in
    its column's format as Excel gives it. ``delete_rows`` are (sheet, row key). ``sort`` is (sheet, heading)
    or (sheet, heading, reverse), sorting whole rows. ``reorder_columns`` reverses the columns of every data
    sheet and ``extra_column`` adds one of the owner's own. ``drop_custom_props`` and ``drop_validations``
    are what another program may lose. The file is saved with a shared strings table, as Excel saves it."""
    from io import BytesIO

    from openpyxl import load_workbook
    from openpyxl.packaging.custom import CustomPropertyList
    from openpyxl.utils import get_column_letter

    from .workbook import SHEETS, SHEET_BY_NAME, fold

    book = load_workbook(BytesIO(data))
    data_sheets = [book[sheet.name] for sheet in SHEETS if sheet.name in book.sheetnames]

    def headings(ws):
        return {cell.value: cell.column for cell in ws[1] if cell.value is not None}

    def key_column(ws):
        sheet = SHEET_BY_NAME[fold(ws.title)]
        return headings(ws)[next(c.heading for c in sheet.columns if c.key == sheet.key)]

    def find_row(ws, key):
        column = key_column(ws)
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row=row, column=column).value) == str(key):
                return row
        raise KeyError(f"{ws.title} has no row {key}")

    def put(cell, value):
        if not cell.has_style:
            # Excel gives a new cell its column's format.
            cell.number_format = cell.parent.column_dimensions[get_column_letter(cell.column)].number_format
        if isinstance(value, str):
            value = excel_typed(value, cell.number_format)
        if isinstance(value, Pasted):
            cell.value = value.value
        else:
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"   # typed into a text cell: = stays text

    # Rows are found as downloaded, before any edit, so edits that swap ids find the rows they meant.
    cells = [(book[sheet_name].cell(row=find_row(book[sheet_name], key), column=headings(book[sheet_name])[heading]),
              value) for sheet_name, key, heading, value in edits]
    for cell, value in cells:
        put(cell, value)
    for sheet_name, key in delete_rows:
        ws = book[sheet_name]
        ws.delete_rows(find_row(ws, key))
    for sheet_name, values in insert_rows:
        ws = book[sheet_name]
        row = ws.max_row + 1
        for heading, value in values.items():
            put(ws.cell(row=row, column=headings(ws)[heading]), value)

    def grid(ws):
        return [[(cell.value, cell.number_format, cell.data_type) for cell in row]
                for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)]

    def rewrite(ws, rows):
        for r, row in enumerate(rows, start=1):
            for c, (value, number_format, data_type) in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c)
                cell.hyperlink = None
                cell.value = value
                cell.number_format = number_format
                if value is not None:
                    cell.data_type = data_type

    if sort:
        sheet_name, heading, *reverse = sort
        ws = book[sheet_name]
        rows = grid(ws)
        at = headings(ws)[heading] - 1
        body = sorted(rows[1:], key=lambda row: (row[at][0] is None, str(row[at][0] or "")),
                      reverse=bool(reverse and reverse[0]))
        rewrite(ws, [rows[0], *body])
    if reorder_columns:
        for ws in data_sheets:
            rewrite(ws, [list(reversed(row)) for row in grid(ws)])
    if extra_column:
        for ws in data_sheets:
            column = ws.max_column + 1
            ws.cell(row=1, column=column, value="My notes")
            ws.cell(row=2, column=column, value="check this one")
    if drop_custom_props:
        book.custom_doc_props = CustomPropertyList()
    if drop_validations:
        for ws in book.worksheets:
            ws.data_validations.dataValidation = []
    out = BytesIO()
    book.save(out)
    return shared_strings(out.getvalue())


def shared_strings(data):
    """``data`` with every inline string moved into one shared strings table, which is how Excel writes text
    (openpyxl writes it inline)."""
    import re
    import zipfile
    from io import BytesIO

    inline = re.compile(r'<c ([^>]*?)t="inlineStr"([^>]*)><is><t(?: xml:space="preserve")?>(.*?)</t></is></c>', re.S)
    empty = re.compile(r'<c ([^>]*?)t="inlineStr"([^>]*)><is><t(?: xml:space="preserve")?\s*/></is></c>')
    strings, index = [], {}

    def shared(match):
        text = match.group(3)
        if text not in index:
            index[text] = len(strings)
            strings.append(text)
        return f'<c {match.group(1)}t="s"{match.group(2)}><v>{index[text]}</v></c>'

    source = zipfile.ZipFile(BytesIO(data))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            body = source.read(info.filename)
            if info.filename.startswith("xl/worksheets/sheet"):
                text = empty.sub(lambda m: f"<c {m.group(1)}{m.group(2).strip()}/>".replace("  ", " "), body.decode())
                body = inline.sub(shared, text).encode()
            elif info.filename == "[Content_Types].xml":
                body = body.decode().replace("</Types>", (
                    '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-'
                    'officedocument.spreadsheetml.sharedStrings+xml"/></Types>')).encode()
            elif info.filename == "xl/_rels/workbook.xml.rels":
                body = body.decode().replace("</Relationships>", (
                    '<Relationship Id="rIdShared" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    'relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>')).encode()
            target.writestr(info, body)
        items = "".join(f'<si><t xml:space="preserve">{text}</t></si>' for text in strings)
        target.writestr("xl/sharedStrings.xml", (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            f'count="{len(strings)}" uniqueCount="{len(strings)}">{items}</sst>'))
    return out.getvalue()
