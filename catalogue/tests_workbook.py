"""The owner's spreadsheet download: the sheets, what each row says, cells Excel keeps as written, the values
an upload is compared with, and nothing personal in the file."""

import gzip
import json
import os
import tempfile
import zipfile
from datetime import date, datetime, timezone as dt_timezone
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.management import CommandError, call_command
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook

from . import data_jobs, workbook
from .models import (
    CheckAnswer, ClaudeJudge, DailySearch, DailyVisitor, DataExport, ImportRun, Listing, Product, ProductAlias,
    Retailer, ShopProduct, StockAlert,
)
from .testing import make_game, make_listing, make_product, make_retailer, make_set

Kind = CheckAnswer.Kind
Status = ShopProduct.Status

PRODUCT_HEADINGS = [
    "Product id", "Address", "Game", "Name", "Type", "Set", "Set code", "Language", "Barcode", "Release date",
    "Show on site", "Same as product id", "Not the same as", "Look-alike group", "Kept apart from", "Shops",
    "In stock at", "Cheapest", "Added", "Notes",
]
LISTING_HEADINGS = [
    "Listing id", "Shop", "Shop's title", "Link", "Product id", "Product", "Product shown", "Price", "Delivery",
    "Stock", "Last checked", "Price check", "Shop's barcode", "Hidden by", "Show on site", "Not this product",
]
SHOP_ITEM_HEADINGS = [
    "Shop item id", "Shop", "Shop's title", "Link", "Shop's barcode", "Price seen", "Stock seen", "Last seen",
    "Found by", "Status", "Suggested id", "Suggested product", "Match %", "Already shown there",
    "Link to product id", "Decision",
]
LEFT_OUT_HEADINGS = ["Shop", "Shop's title", "Link", "Barcode given", "Why left out", "Link to product id", "Decision"]
SET_HEADINGS = ["Set id", "Address", "Game", "Name", "Code", "Release date", "Date from", "Products"]
EDITABLE = {
    "Products": {"Name", "Type", "Set", "Language", "Barcode", "Release date", "Show on site", "Same as product id",
                 "Not the same as"},
    "Listings": {"Product id", "Show on site", "Not this product"},
    "Shop items": {"Link to product id", "Decision"},
    "Left out": {"Link to product id", "Decision"},
    "Sets": {"Name", "Code", "Release date"},
}

# Left-out lines as a read keeps them: the first two look sealed, the last two do not.
ILLUSTRATION = "One Piece Card Game - Illustration Box (IB-06) [no barcode] https://plain.example/products/ib-06"
SURGING = "Pokemon Surging Sparks Booster Box [0820650851230] https://plain.example/products/ssb"
CASE = "Disney Lorcana Azurite Sea Booster Box Case (4 Boxes) [no barcode] https://plain.example/products/case"
BOARD_GAME = "Ticket to Ride Board Game [no barcode] https://plain.example/products/ttr"


class Base(TestCase):
    def setUp(self):
        cache.clear()
        self.game = make_game()
        self.set = make_set(self.game, name="Surging Sparks", slug="surging-sparks", code="SSP",
                            release_date=date(2024, 11, 8))
        self.box = make_product(self.set, name="Surging Sparks Booster Box", slug="ssp-box",
                                product_type="booster_box", ean="0123456789012", release_date=date(2026, 11, 15))
        self.shop = make_retailer("Solent Cards")

    def export(self, **kwargs):
        export = DataExport.objects.create(**kwargs)
        data, base, counts = workbook.build_export(export)
        return export, data, base, counts

    def rows(self, data, sheet):
        """The sheet's data rows as {heading: value}, as openpyxl reads them back."""
        ws = load_workbook(BytesIO(data))[sheet]
        lines = list(ws.iter_rows(values_only=True))
        return [dict(zip(lines[0], line)) for line in lines[1:]]

    def by_id(self, data, sheet, heading):
        return {row[heading]: row for row in self.rows(data, sheet)}


class SheetTests(Base):
    def test_the_sheets_their_headings_and_which_he_may_change(self):
        export, data, _base, counts = self.export()
        book = load_workbook(BytesIO(data))
        self.assertEqual(book.sheetnames, ["Read me", "Products", "Listings", "Shop items", "Sets", "Lists"])
        self.assertEqual(book["Lists"].sheet_state, "hidden")
        expected = {"Products": PRODUCT_HEADINGS, "Listings": LISTING_HEADINGS, "Shop items": SHOP_ITEM_HEADINGS,
                    "Sets": SET_HEADINGS}
        for name, headings in expected.items():
            ws = book[name]
            heading_cells = next(ws.iter_rows(min_row=1, max_row=1))
            self.assertEqual([cell.value for cell in heading_cells], headings)
            self.assertEqual(ws.freeze_panes, "B2")
            self.assertTrue(ws.auto_filter.ref.startswith("A1:"))
            for cell in heading_cells:
                with self.subTest(sheet=name, heading=cell.value):
                    self.assertTrue(cell.font.b)
                    orange = cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb.endswith(workbook.EDITABLE_FILL)
                    self.assertEqual(orange, cell.value in EDITABLE[name])
            # The grey of the read-only columns is one rule a column, not a fill on every cell.
            grey = {str(rule.sqref).split(":")[0].rstrip("1") for rule in ws.conditional_formatting}
            letters = {cell.column_letter for cell in heading_cells if cell.value not in EDITABLE[name]}
            self.assertEqual(grey, letters)
        self.assertEqual(counts, {"Products": 1, "Listings": 0, "Shop items": 0, "Sets": 1})

    def test_the_spec_and_the_sheets_agree(self):
        for sheet in workbook.SHEETS:
            keys = [column.key for column in sheet.columns]
            self.assertEqual(len(keys), len(set(keys)), sheet.name)
            self.assertIn(sheet.key, keys)
            if sheet.integrity:
                self.assertIn(sheet.integrity, keys)
            for column in sheet.columns:
                self.assertIn(column.mark, ("E", "R", "A"))
                if column.dropdown:
                    self.assertIn(column.dropdown, workbook.LIST_NAMES)

    def test_the_file_number_is_in_the_properties_and_on_read_me(self):
        export, data, _base, _counts = self.export()
        book = load_workbook(BytesIO(data))
        props = {prop.name: prop.value for prop in book.custom_doc_props}
        self.assertEqual(props["RipRaptor file"], export.reference)
        self.assertRegex(export.reference, rf"^RR-{export.pk}-[0-9a-f]{{12}}$")
        readme = book["Read me"]
        self.assertEqual(readme["A1"].value, "RipRaptor data")
        self.assertEqual((readme["A2"].value, readme["B2"].value), ("File reference", export.reference))
        self.assertEqual(readme["A3"].value, "Downloaded")
        self.assertIsInstance(readme["B3"].value, datetime)
        self.assertEqual(readme["B4"].value, "Products 1, Listings 0, Shop items 0, Sets 1")
        lines = [row[0] for row in readme.iter_rows(min_row=6, values_only=True)]
        self.assertEqual(lines, list(workbook.README_LINES))
        # The upload page finds the number without openpyxl.
        self.assertEqual(workbook.quick_reference(zipfile.ZipFile(BytesIO(data))), (export.pk, export.token))

    def test_read_me_says_it_plainly(self):
        for line in workbook.README_LINES:
            with self.subTest(line=line):
                self.assertNotIn("—", line)
                self.assertNotIn("!", line)
                for phrase in ("learn more", "get started", "unlock", "seamless", "all in one place", "effortless"):
                    self.assertNotIn(phrase, line.lower())

    def test_dropdowns_read_the_hidden_lists_and_only_warn(self):
        _export, data, _base, _counts = self.export()
        book = load_workbook(BytesIO(data))
        lists = list(book["Lists"].iter_rows(values_only=True))
        self.assertEqual([row[0] for row in lists[:2]], ["Yes", "No"])
        self.assertIn("Booster box", [row[1] for row in lists])
        self.assertIn("Pokémon", [row[2] for row in lists])
        self.assertIn("Surging Sparks", [row[3] for row in lists])
        self.assertEqual([row[4] for row in lists][:2], ["English", "Japanese"])
        self.assertEqual([row[5] for row in lists][:3], ["Link", "Not ours", "Look again"])
        checks = {str(dv.sqref): dv for dv in book["Products"].data_validations.dataValidation}
        types = checks["E2:E2002"]
        self.assertEqual(types.type, "list")
        self.assertTrue(types.formula1.startswith("Lists!$B$1:$B$"))
        self.assertEqual(types.errorStyle, "warning")
        self.assertEqual(checks["K2:K2002"].formula1, "Lists!$A$1:$A$2")


class CellTests(Base):
    def test_barcodes_stay_text_and_dates_stay_dates(self):
        make_product(self.set, name="Surging Sparks Elite Trainer Box", ean="5053222123456")
        _export, data, base, _counts = self.export()
        rows = {row["Name"]: row for row in self.rows(data, "Products")}
        self.assertEqual(rows["Surging Sparks Booster Box"]["Barcode"], "0123456789012")
        self.assertEqual(rows["Surging Sparks Elite Trainer Box"]["Barcode"], "5053222123456")
        released = rows["Surging Sparks Booster Box"]["Release date"]
        self.assertEqual(released, datetime(2026, 11, 15))
        ws = load_workbook(BytesIO(data))["Products"]
        self.assertEqual(ws["I2"].number_format, "@")
        self.assertEqual(ws["J2"].number_format, "dd/mm/yyyy")
        self.assertTrue(ws["J2"].is_date)
        # Rows he adds keep what he types as text in the editable columns.
        self.assertEqual(ws.column_dimensions["I"].number_format, "@")
        self.assertEqual(ws.column_dimensions["L"].number_format, "@")
        self.assertEqual(base["sheets"]["Products"][str(self.box.pk)]["barcode"], "0123456789012")
        self.assertEqual(base["sheets"]["Products"][str(self.box.pk)]["release_date"], "2026-11-15")

    def test_times_are_london_times(self):
        checked = datetime(2026, 7, 1, 11, 0, tzinfo=dt_timezone.utc)
        listing = make_listing(self.box, self.shop)
        Listing.objects.filter(pk=listing.pk).update(last_checked=checked)
        _export, data, _base, _counts = self.export()
        self.assertEqual(self.by_id(data, "Listings", "Listing id")[listing.pk]["Last checked"], datetime(2026, 7, 1, 12, 0))

    def test_a_title_that_looks_like_a_formula_is_written_as_text(self):
        listing = make_listing(self.box, self.shop, title='=HYPERLINK("http://evil.example","x")')
        ShopProduct.objects.create(retailer=self.shop, title="=1+1", url="https://solent.example/products/sum",
                                   price=Decimal("9"), status=Status.REVIEW)
        Product.objects.filter(pk=self.box.pk).update(name="@SUM(1)")
        _export, data, _base, _counts = self.export()
        with zipfile.ZipFile(BytesIO(data)) as archive:
            for name in archive.namelist():
                if name.startswith("xl/worksheets/sheet"):
                    xml = archive.read(name).decode()
                    self.assertNotIn("<f>", xml, name)
                    self.assertNotIn("<f ", xml, name)
        self.assertEqual(self.by_id(data, "Listings", "Listing id")[listing.pk]["Shop's title"],
                         '=HYPERLINK("http://evil.example","x")')

    def test_a_character_excel_cannot_hold_is_dropped(self):
        listing = make_listing(self.box, self.shop, title="Booster\x0bBox\x01")
        _export, data, _base, _counts = self.export()
        self.assertEqual(self.by_id(data, "Listings", "Listing id")[listing.pk]["Shop's title"], "BoosterBox")

    def test_marketplace_links_lose_their_query_string(self):
        ebay = make_retailer("eBay", source_type=Retailer.Source.EBAY)
        amazon = make_retailer("Amazon", source_type=Retailer.Source.AMAZON)
        on_ebay = make_listing(self.box, ebay, url="https://www.ebay.co.uk/itm/123?mkcid=1&campid=5338")
        on_amazon = make_listing(self.box, amazon, url="https://www.amazon.co.uk/dp/B0TEST?tag=ripraptor-21")
        shop = make_listing(self.box, self.shop, url="https://solent.example/products/box?variant=7")
        _export, data, base, _counts = self.export()
        links = {pk: row["Link"] for pk, row in self.by_id(data, "Listings", "Listing id").items()}
        self.assertEqual(links[on_ebay.pk], "https://www.ebay.co.uk/itm/123")
        self.assertEqual(links[on_amazon.pk], "https://www.amazon.co.uk/dp/B0TEST")
        self.assertEqual(links[shop.pk], "https://solent.example/products/box?variant=7")
        self.assertEqual(base["sheets"]["Listings"][str(on_ebay.pk)]["link"], "https://www.ebay.co.uk/itm/123")
        with zipfile.ZipFile(BytesIO(data)) as archive:
            for name in archive.namelist():
                text = archive.read(name).decode("utf-8", "replace")
                self.assertNotIn("campid", text, name)
                self.assertNotIn("ripraptor-21", text, name)
        ws = load_workbook(BytesIO(data))["Listings"]
        self.assertEqual(ws["D2"].hyperlink.target, ws["D2"].value)


class ProductSheetTests(Base):
    def test_a_product_row_and_its_base(self):
        listing = make_listing(self.box, self.shop, price="90.00", delivery="0.00")
        _export, data, base, _counts = self.export()
        row = self.by_id(data, "Products", "Product id")[self.box.pk]
        self.assertEqual(row["Address"], f"https://ripraptor.com/products/{self.box.slug}/")
        self.assertEqual((row["Game"], row["Name"], row["Type"], row["Set"], row["Set code"], row["Language"]),
                         ("Pokémon", "Surging Sparks Booster Box", "Booster box", "Surging Sparks", "SSP", "English"))
        self.assertEqual((row["Show on site"], row["Shops"], row["In stock at"], row["Cheapest"]), ("Yes", 1, 1, 90))
        self.assertIsNone(row["Same as product id"])
        self.assertIsNone(row["Look-alike group"])
        self.assertEqual(base["v"], 1)
        self.assertEqual(base["sheets"]["Products"][str(self.box.pk)], {
            "address": row["Address"], "name": "Surging Sparks Booster Box", "type": "booster_box",
            "set": self.set.pk, "language": "", "barcode": "0123456789012", "release_date": "2026-11-15",
            "show": True,
        })
        self.assertEqual(base["sheets"]["Listings"][str(listing.pk)],
                         {"link": listing.url, "product": self.box.pk, "show": True})

    def test_a_language_and_a_hidden_product(self):
        japanese = make_product(self.set, name="Surging Sparks Booster Box (Japanese)", is_active=False)
        _export, data, base, _counts = self.export()
        row = self.by_id(data, "Products", "Product id")[japanese.pk]
        self.assertEqual((row["Language"], row["Show on site"]), ("Japanese", "No"))
        self.assertEqual(base["sheets"]["Products"][str(japanese.pk)]["language"], "ja")
        self.assertFalse(base["sheets"]["Products"][str(japanese.pk)]["show"])

    def test_merged_away_products_are_left_out_and_their_listings_say_where_they_went(self):
        gone = make_product(self.set, name="Surging Sparks Booster Display", slug="ssp-display", is_active=False)
        ProductAlias.objects.create(slug=gone.slug, product=self.box)
        hidden = make_product(self.set, name="Surging Sparks Mini Tin", is_active=False)
        left = make_listing(gone, self.shop, url="https://solent.example/products/display")
        _export, data, base, _counts = self.export()
        products = self.by_id(data, "Products", "Product id")
        self.assertNotIn(gone.pk, products)
        self.assertIn(hidden.pk, products)
        self.assertNotIn(str(gone.pk), base["sheets"]["Products"])
        self.assertEqual(self.by_id(data, "Listings", "Listing id")[left.pk]["Product shown"], f"Merged into {self.box.pk}")

    def test_look_alike_groups_and_pairs_kept_apart(self):
        loose = make_product(self.set, name="Journey Together Elite Trainer Box")
        loose_too = make_product(self.set, name="Journey Together Elite Trainer Box (Exclusive)")
        barcode = make_product(self.set, name="Mega Evolution Booster Bundle", ean="0820650851230")
        barcode_too = make_product(self.set, name="Some Other Tin", product_type="tin", ean="820650851230")
        typed = make_product(self.set, name="Pokemon Scarlet & Violet Surging Sparks Booster Box", product_type="bundle")
        alone = make_product(self.set, name="Black Bolt Binder Collection", product_type="collection_box")
        CheckAnswer.objects.create(kind=Kind.APART, by_owner=True, what="Not the same", product=loose, other=loose_too)
        CheckAnswer.objects.create(kind=Kind.APART, by_owner=True, what="Not the same", product=alone, other=loose,
                                   undone_at=timezone.now())
        _export, data, _base, _counts = self.export()
        rows = self.by_id(data, "Products", "Product id")
        group = {pk: row["Look-alike group"] for pk, row in rows.items()}
        self.assertIsNotNone(group[loose.pk])
        self.assertEqual(group[loose.pk], group[loose_too.pk])
        self.assertIsNotNone(group[barcode.pk])
        self.assertEqual(group[barcode.pk], group[barcode_too.pk])
        # The type is no part of the key: the box filed as a bundle meets the box.
        self.assertIsNotNone(group[typed.pk])
        self.assertEqual(group[typed.pk], group[self.box.pk])
        self.assertIsNone(group[alone.pk])
        self.assertEqual(len({group[loose.pk], group[barcode.pk], group[self.box.pk]}), 3)
        self.assertEqual(rows[loose.pk]["Kept apart from"], str(loose_too.pk))
        self.assertEqual(rows[loose_too.pk]["Kept apart from"], str(loose.pk))
        self.assertIsNone(rows[alone.pk]["Kept apart from"])

    def test_look_alike_groups_number_in_order_and_join_a_shared_code(self):
        groups = workbook.lookalike_groups([
            (1, 1, "pokemon", "Pursue Glory Together csv10c Booster Box", "zh-hans", ""),
            (2, 1, "pokemon", "Lonely Tin", "", ""),
            (3, 1, "pokemon", "Chasing Glory Together CSV10C Booster Box", "zh", ""),
            (4, 2, "pokemon", "Chasing Glory Together CSV10C Booster Box", "zh", ""),   # another game
            (5, 1, "pokemon", "Surging Sparks Booster Box", "ja", ""),
            (6, 1, "pokemon", "Surging Sparks Booster Box", "", ""),   # another language
            (7, 3, "magic-the-gathering", "Foundations Play Booster Box", "", "5010996000001"),
            (8, 4, "lorcana", "Unrelated Name", "", "05010996000001"),   # one barcode, Excel or not
        ])
        self.assertEqual(groups, {1: 1, 3: 1, 7: 2, 8: 2})

    def test_notes_say_what_the_hourly_tidy_will_do(self):
        from django.contrib.admin.models import CHANGE, LogEntry

        owner = get_user_model().objects.create_superuser("ben", "ben@example.com", "pw")
        LogEntry.objects.log_actions(owner.pk, Product.objects.filter(pk=self.box.pk), CHANGE,
                                     change_message=[{"changed": {"fields": ["Product type"]}}], single_object=True)
        doomed = make_product(self.set, name="Surging Sparks Card Sleeves Box")
        _export, data, _base, _counts = self.export()
        rows = self.by_id(data, "Products", "Product id")
        self.assertEqual(rows[self.box.pk]["Notes"], "Type set by you")
        self.assertEqual(rows[doomed.pk]["Notes"], "No shop and no release date: the hourly tidy removes it")


class ListingSheetTests(Base):
    def test_a_listing_row(self):
        listing = make_listing(self.box, self.shop, price="84.50", delivery="3.99", title="SSP Booster Box",
                               shop_ean="0820650851230", availability=Listing.Availability.PREORDER)
        unknown = make_listing(make_product(self.set, name="Surging Sparks Tin", product_type="tin"),
                               self.shop, delivery_known=False, sanity=Listing.Sanity.EXCLUDED,
                               sanity_reason="far below every other shop")
        _export, data, _base, _counts = self.export()
        rows = self.by_id(data, "Listings", "Listing id")
        row = rows[listing.pk]
        self.assertEqual((row["Shop"], row["Shop's title"], row["Product id"], row["Product"], row["Product shown"]),
                         ("Solent Cards", "SSP Booster Box", self.box.pk, "Surging Sparks Booster Box", "Yes"))
        self.assertEqual((row["Price"], row["Delivery"], row["Stock"], row["Price check"], row["Shop's barcode"]),
                         (84.5, 3.99, "Pre-order", "OK", "0820650851230"))
        self.assertEqual((row["Show on site"], row["Hidden by"], row["Not this product"]), ("Yes", None, None))
        self.assertIsNone(rows[unknown.pk]["Delivery"])
        self.assertEqual(rows[unknown.pk]["Price check"], "Kept out: far below every other shop")

    def test_hidden_by_says_who_hid_it(self):
        from .models import ClaudeAsk

        products = [make_product(self.set, name=f"Surging Sparks Pack {n}", product_type="booster_pack") for n in range(4)]
        listings = [make_listing(product, self.shop, is_active=False) for product in products]
        CheckAnswer.objects.create(kind=Kind.HIDE, by_owner=True, what="Hid", listing=listings[0])
        CheckAnswer.objects.create(kind=Kind.HIDE, what="Hid", listing=listings[1])
        ask = ClaudeAsk.objects.create(kind=ClaudeAsk.Kind.OFFER, row_key=f"offer:{listings[2].pk}", fingerprint="f",
                                       model_asked="m", effort="low")
        CheckAnswer.objects.create(kind=Kind.HIDE, what="Hid", listing=listings[2], ask=ask)
        # Undone, and a hide on a listing shown again: neither names anyone.
        CheckAnswer.objects.create(kind=Kind.HIDE, what="Hid", listing=listings[3], undone_at=timezone.now())
        _export, data, base, _counts = self.export()
        rows = self.by_id(data, "Listings", "Listing id")
        self.assertEqual([rows[listing.pk]["Hidden by"] for listing in listings], ["you", "autopilot", "Claude", None])
        self.assertEqual([rows[listing.pk]["Show on site"] for listing in listings], ["No"] * 4)
        self.assertFalse(base["sheets"]["Listings"][str(listings[0].pk)]["show"])


class ShopItemTests(Base):
    def item(self, title, status, suggested=None, price=Decimal("80.00"), slug=None):
        slug = slug or title.lower().replace(" ", "-")
        return ShopProduct.objects.create(retailer=self.shop, title=title, url=f"https://solent.example/products/{slug}",
                                          price=price, status=status, suggested=suggested, confidence=70)

    def test_waiting_ignored_and_linked_items_the_site_does_not_show(self):
        other = make_product(self.set, name="Surging Sparks Elite Trainer Box")
        make_listing(other, self.shop, url="https://solent.example/products/etb-page-one", title="SSP ETB")
        tin = make_product(self.set, name="Surging Sparks Tin", product_type="tin")
        make_listing(tin, self.shop, url="https://solent.example/products/tin")
        no_product = self.item("Surging Sparks Display", Status.LINKED, suggested=None)
        no_listing = self.item("Surging Sparks Booster Box", Status.LINKED, suggested=self.box)
        waiting = self.item("Surging Sparks Box Unpriced", Status.LINKED, suggested=self.box, price=None)
        review = self.item("Surging Sparks ETB Pokemon Center", Status.REVIEW, suggested=other)
        ignored = self.item("Surging Sparks Sleeves", Status.IGNORED)
        shown = self.item("Surging Sparks Tin", Status.LINKED, suggested=tin, slug="tin")
        merged = make_product(self.set, name="Surging Sparks Bundle", product_type="bundle", is_active=False)
        on_merged = self.item("Surging Sparks Bundle", Status.LINKED, suggested=merged)
        second_page = self.item("Surging Sparks ETB Second Page", Status.LINKED, suggested=other)
        _export, data, base, counts = self.export()
        rows = self.by_id(data, "Shop items", "Shop item id")
        self.assertNotIn(shown.pk, rows)
        self.assertEqual(counts["Shop items"], 7)
        self.assertEqual(rows[second_page.pk]["Status"],
                         "Linked, not shown: this shop's other page is already shown on the product")
        self.assertEqual(rows[second_page.pk]["Already shown there"], "SSP ETB")
        self.assertEqual(rows[no_product.pk]["Status"], workbook.LINKED_GONE)
        self.assertEqual(rows[on_merged.pk]["Status"], workbook.LINKED_GONE)
        self.assertEqual(rows[no_listing.pk]["Status"], "Linked, not shown: no listing")
        self.assertEqual(rows[waiting.pk]["Status"], "Linked, waiting for the shop's next read")
        self.assertEqual(rows[ignored.pk]["Status"], "Not ours")
        self.assertEqual(rows[review.pk]["Status"],
                         "Waiting for you: this shop's other page is already shown on the suggested product")
        self.assertEqual(rows[review.pk]["Already shown there"], "SSP ETB")
        self.assertEqual((rows[review.pk]["Suggested id"], rows[review.pk]["Suggested product"], rows[review.pk]["Match %"]),
                         (other.pk, "Surging Sparks Elite Trainer Box", 70))
        # Waiting first, then linked, then not ours.
        self.assertEqual(list(rows)[0], review.pk)
        self.assertEqual(list(rows)[-1], ignored.pk)
        self.assertEqual(base["sheets"]["Shop items"][str(review.pk)],
                         {"link": review.url, "status": "review", "suggested": other.pk})


class LeftOutTests(Base):
    def setUp(self):
        super().setUp()
        ImportRun.objects.create(retailer=self.shop, finished_at=timezone.now(), unmatched="\r\n".join([
            ILLUSTRATION, SURGING, CASE, BOARD_GAME,
            "Pokemon Draft Night Box [no barcode] https://plain.example/products/draft",
            "Pokemon Prismatic Booster Bundle [no barcode] https://plain.example/products/bundle",
            "Pokemon Mega Brave Booster Box -> maybe Mega Brave Booster Box (80%)",
        ]))
        ShopProduct.objects.create(retailer=self.shop, title="Draft", url="https://plain.example/products/draft")
        make_listing(self.box, self.shop, url="https://www.plain.example/products/bundle/")

    def test_only_with_a_shop_chosen_and_only_what_the_site_does_not_know(self):
        from .insights import left_out_lines

        # Four lines look sealed; the draft night has a shop item row and the bundle a listing at this shop.
        self.assertEqual(len(list(left_out_lines(ImportRun.objects.get().unmatched))), 4)
        _export, data, _base, counts = self.export()
        self.assertNotIn("Left out", load_workbook(BytesIO(data)).sheetnames)
        self.assertNotIn("Left out", counts)
        _export, data, base, counts = self.export(left_out_retailer=self.shop)
        book = load_workbook(BytesIO(data))
        self.assertEqual(book.sheetnames, ["Read me", "Products", "Listings", "Shop items", "Left out", "Sets", "Lists"])
        self.assertEqual([cell.value for cell in next(book["Left out"].iter_rows(max_row=1))], LEFT_OUT_HEADINGS)
        rows = self.rows(data, "Left out")
        self.assertEqual([(r["Shop"], r["Link"], r["Barcode given"]) for r in rows], [
            ("Solent Cards", "https://plain.example/products/ib-06", None),
            ("Solent Cards", "https://plain.example/products/ssb", "0820650851230"),
        ])
        self.assertEqual(rows[0]["Why left out"], "a kind of product the rules do not know")
        self.assertEqual(counts["Left out"], 2)
        key = f"{self.shop.pk}|plain.example/products/ssb"
        self.assertEqual(base["sheets"]["Left out"][key], {
            "link": "https://plain.example/products/ssb", "title": "Pokemon Surging Sparks Booster Box",
            "barcode": "0820650851230"})

    def test_only_the_latest_good_whole_read(self):
        ImportRun.objects.create(retailer=self.shop, finished_at=timezone.now(), error="boom", unmatched=BOARD_GAME)
        ImportRun.objects.create(retailer=self.shop, finished_at=timezone.now(), note="pre-order pulse",
                                 unmatched=SURGING)
        _export, data, _base, counts = self.export(left_out_retailer=self.shop)
        self.assertEqual(counts["Left out"], 2)
        ImportRun.objects.create(retailer=self.shop, finished_at=timezone.now(), unmatched=SURGING)
        _export, data, _base, counts = self.export(left_out_retailer=self.shop)
        self.assertEqual([r["Link"] for r in self.rows(data, "Left out")], ["https://plain.example/products/ssb"])


class SetSheetTests(Base):
    def test_a_set_row_and_its_base(self):
        from .releases import OWNER

        mine = make_set(self.game, name="Black Bolt", slug="black-bolt", code="BLK", release_date=date(2025, 7, 18),
                        release_date_source=OWNER)
        _export, data, base, _counts = self.export()
        rows = self.by_id(data, "Sets", "Set id")
        self.assertEqual(list(rows), [mine.pk, self.set.pk])   # newest first
        self.assertEqual((rows[mine.pk]["Name"], rows[mine.pk]["Code"], rows[mine.pk]["Date from"],
                          rows[mine.pk]["Products"]), ("Black Bolt", "BLK", "you", 0))
        self.assertEqual(rows[self.set.pk]["Products"], 1)
        self.assertEqual(rows[mine.pk]["Release date"], datetime(2025, 7, 18))
        self.assertEqual(base["sheets"]["Sets"][str(mine.pk)], {
            "address": "https://ripraptor.com/games/pokemon/black-bolt/", "name": "Black Bolt", "code": "BLK",
            "release_date": "2025-07-18"})


class PrivacyTests(Base):
    def test_nothing_personal_or_secret_goes_in_the_file(self):
        from inbox.models import Conversation, Message

        owner = get_user_model().objects.create_superuser("ben", "owner.private@example.com", "pw")
        StockAlert.objects.create(product=self.box, email="alert.person@example.org")
        thread = Conversation.objects.create(name="Visitor Name", email="visitor.private@example.net")
        Message.objects.create(conversation=thread, body="my private message xyzzy")
        DailySearch.objects.create(date=date.today(), query="secret search phrase")
        DailyVisitor.objects.create(date=date.today(), token="visitortoken1234")
        Retailer.objects.filter(pk=self.shop.pk).update(
            affiliate_url_template="https://aff.example/click?id=SECRETAFF&url={url}",
            session_url="https://solent.example/session-secret",
            source_url="https://feed.example/products.csv?apikey=SECRETAPIKEY",
        )
        make_listing(self.box, self.shop)
        ClaudeJudge.objects.create(pk=1, key_hint="Zq9X")
        ImportRun.objects.create(retailer=self.shop, finished_at=timezone.now(), unmatched=SURGING,
                                 log="GET https://feed.example/products.csv?apikey=SECRETAPIKEY")
        export = DataExport.objects.create(created_by=owner, left_out_retailer=self.shop)
        data = data_jobs.build_export(export)
        export.refresh_from_db()
        secrets = ["alert.person@example.org", "visitor.private@example.net", "Visitor Name", "xyzzy",
                   "secret search phrase", "visitortoken1234", "SECRETAFF", "aff.example", "session-secret",
                   "SECRETAPIKEY", "feed.example", "Zq9X", "owner.private@example.com"]
        texts = {}
        with zipfile.ZipFile(BytesIO(data)) as archive:
            for name in archive.namelist():
                texts[name] = archive.read(name).decode("utf-8", "replace")
        texts["base"] = gzip.decompress(bytes(export.base)).decode()
        self.assertIn(self.box.name, texts["base"])
        for name, text in texts.items():
            for secret in secrets:
                with self.subTest(member=name, secret=secret):
                    self.assertNotIn(secret, text)


class QueryCountTests(Base):
    def add(self, n, start):
        for i in range(start, start + n):
            product = make_product(self.set, name=f"Surging Sparks Thing {i}", slug=f"thing-{i}")
            make_listing(product, self.shop, url=f"https://solent.example/products/thing-{i}")
            make_listing(product, make_retailer(f"Shop {i}"))
            ShopProduct.objects.create(retailer=self.shop, title=f"Thing {i} Box", url=f"https://solent.example/p/{i}",
                                       status=Status.REVIEW, suggested=product)
        CheckAnswer.objects.create(kind=Kind.HIDE, what="Hid", listing=Listing.objects.last())
        CheckAnswer.objects.create(kind=Kind.APART, by_owner=True, what="Apart", product=self.box,
                                   other=Product.objects.last())

    def test_the_queries_do_not_grow_with_the_catalogue(self):
        ImportRun.objects.create(retailer=self.shop, finished_at=timezone.now(), unmatched=SURGING)
        self.add(3, 0)
        export = DataExport.objects.create(left_out_retailer=self.shop)
        workbook.build_export(export)   # the first build looks up content types; later ones do not
        with CaptureQueriesContext(connection) as small:
            workbook.build_export(export)
        self.add(27, 3)
        with self.assertNumQueries(len(small)):
            data, _base, counts = workbook.build_export(export)
        self.assertEqual(counts["Products"], 31)
        self.assertEqual(counts["Listings"], 60)
        self.assertEqual(counts["Shop items"], 30)
        self.assertLessEqual(len(small), 12)


class CommandTests(Base):
    def test_export_writes_the_file_and_keeps_it_as_a_download(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "data.xlsx"
            out = StringIO()
            call_command("data_files", "--export", str(path), "--left-out", self.shop.slug, stdout=out)
            data = path.read_bytes()
        export = DataExport.objects.get()
        self.assertEqual(export.status, DataExport.Status.READY)
        self.assertEqual(bytes(export.workbook), data)
        self.assertEqual(export.size_bytes, len(data))
        self.assertEqual(export.left_out_retailer, self.shop)
        self.assertEqual(export.counts, {"Products": 1, "Listings": 0, "Shop items": 0, "Left out": 0, "Sets": 1})
        base = data_jobs.unpack_base(export.base)
        self.assertEqual(list(base["sheets"]), ["Products", "Listings", "Shop items", "Left out", "Sets"])
        self.assertIsNotNone(export.finished_at)
        self.assertIn(export.reference, out.getvalue())
        self.assertIn("Products 1", out.getvalue())

    def test_export_needs_a_known_shop_and_a_path(self):
        with self.assertRaisesMessage(CommandError, "No shop has the slug nowhere."):
            call_command("data_files", "--export", "x.xlsx", "--left-out", "nowhere")
        with self.assertRaises(CommandError):
            call_command("data_files")
        self.assertFalse(DataExport.objects.exists())

    def test_a_failed_build_says_so_on_the_row(self):
        export = DataExport.objects.create()
        with mock.patch("catalogue.workbook.build_export", side_effect=RuntimeError("boom")), \
                self.assertLogs("catalogue.data_jobs", "ERROR"):
            self.assertIsNone(data_jobs.build_export(export))
        export.refresh_from_db()
        self.assertEqual(export.status, DataExport.Status.FAILED)
        self.assertEqual(export.error, "The file could not be made (RuntimeError).")
        self.assertIsNone(export.workbook)

    def test_the_file_name_uses_the_london_date(self):
        export = DataExport.objects.create(created_at=datetime(2026, 6, 30, 23, 30, tzinfo=dt_timezone.utc))
        self.assertEqual(export.filename, f"ripraptor-data-2026-07-01-RR{export.pk}.xlsx")
        self.assertEqual(len(export.token), 12)
        self.assertNotEqual(export.token, DataExport.objects.create().token)

    def test_base_packs_and_unpacks(self):
        base = {"v": 1, "sheets": {"Products": {"1": {"name": "Pokémon Box", "release_date": None}}}}
        self.assertEqual(data_jobs.unpack_base(data_jobs.pack_base(base)), base)
        self.assertEqual(json.loads(gzip.decompress(data_jobs.pack_base(base))), base)

    def test_the_web_workers_never_load_openpyxl(self):
        import subprocess
        import sys

        from django.conf import settings

        code = ("import django, sys; django.setup(); import catalogue.workbook, catalogue.data_jobs; "
                "print('openpyxl' in sys.modules)")
        result = subprocess.run([sys.executable, "-c", code], cwd=settings.BASE_DIR, capture_output=True, text=True,
                                env={**os.environ, "DJANGO_SETTINGS_MODULE": "ripraptor.settings"},
                                timeout=60)
        self.assertEqual(result.stdout.strip(), "False", result.stderr)

