"""
The owner's catalogue as an Excel workbook: what a download holds, and the values it was made with.

He downloads the file, tidies it in Excel, saves it and uploads the same file; the upload is compared with
the values written here (the base), never with what the site holds by then, so only what he changed counts.
This module is the one spec for both directions, so the sheets written and the sheets read cannot drift.

openpyxl is imported inside the functions that write or read a workbook, which run in the data_files process,
so a web worker never loads it: the page that takes an upload only needs quick_reference and the spec.
"""

import re
from collections import defaultdict
from tempfile import SpooledTemporaryFile
from typing import NamedTuple

from django.conf import settings
from django.db.models import Case, Count, F, IntegerField, Q, Value, When
from django.urls import reverse
from django.utils import timezone

from . import checks, languages, types
from .importers import ean_key, link_key
from .insights import left_out_lines
from .languages import CHINESE, SIMPLIFIED
from .management.commands.merge_duplicates import chinese_code_key, loose_key
from .models import CheckAnswer, Game, ImportRun, Listing, Product, ProductAlias, ProductSet, ShopProduct
from .releases import OWNER
from .sanity import MARKETPLACES

# Orange headings mark what he may change: --color-brand-soft in web/static/css/tokens.css. The tokens rule is
# for site.css; Excel cannot read a CSS variable, so the value is written here.
EDITABLE_FILL = "FFF1EA"
READ_ONLY_FILL = "EEEEEE"

EDITABLE, READ_ONLY, ACTION = "E", "R", "A"
TEXT_FORMAT = "@"
DATE_FORMAT = "dd/mm/yyyy"
DATETIME_FORMAT = "dd/mm/yyyy hh:mm"
MONEY_FORMAT = '"£"#,##0.00'

# Dropdowns cover the rows written and this many more, for the rows he adds at the bottom.
SPARE_ROWS = 2000
# A shop's left-out list can run to many thousands of lines; the file keeps the first this many.
LEFT_OUT_LIMIT = 5000
# Made in memory up to this size, then in a temporary file.
SPOOL_BYTES = 8 * 1024 * 1024


class Column(NamedTuple):
    heading: str
    key: str
    mark: str    # E editable, R read-only (ignored on upload), A action (exported blank)
    kind: str    # how the cell is written and, on upload, read back
    width: int
    dropdown: str = ""   # the Lists column its dropdown offers


class Sheet(NamedTuple):
    name: str
    key: str          # the column that names the row
    integrity: str    # a read-only column that must still belong to the row's id, or ""
    columns: tuple


E, R, A = EDITABLE, READ_ONLY, ACTION

PRODUCTS = Sheet("Products", "id", "address", (
    Column("Product id", "id", R, "id", 11),
    Column("Address", "address", R, "link", 40),
    # Read for new rows only (new1, new2 and so on): a product never changes game.
    Column("Game", "game", R, "game", 18, "games"),
    Column("Name", "name", E, "name", 60),
    Column("Type", "type", E, "type", 22, "types"),
    Column("Set", "set", E, "set", 28, "sets"),
    Column("Set code", "set_code", R, "text", 10),
    Column("Language", "language", E, "language", 18, "languages"),
    Column("Barcode", "barcode", E, "barcode", 16),
    Column("Release date", "release_date", E, "date", 13),
    Column("Show on site", "show", E, "yes_no", 12, "yes_no"),
    Column("Same as product id", "same_as", A, "ref", 14),
    Column("Not the same as", "not_same", A, "refs", 16),
    Column("Look-alike group", "group", R, "int", 12),
    Column("Kept apart from", "apart", R, "text", 16),
    Column("Shops", "shops", R, "int", 8),
    Column("In stock at", "in_stock", R, "int", 10),
    Column("Cheapest", "cheapest", R, "money", 11),
    Column("Added", "added", R, "datetime", 16),
    Column("Notes", "notes", R, "text", 44),
))

LISTINGS = Sheet("Listings", "id", "link", (
    Column("Listing id", "id", R, "id", 11),
    Column("Shop", "shop", R, "text", 20),
    Column("Shop's title", "title", R, "text", 50),
    Column("Link", "link", R, "link", 40),
    Column("Product id", "product", E, "ref", 11),
    Column("Product", "product_name", R, "text", 45),
    Column("Product shown", "product_shown", R, "text", 15),
    Column("Price", "price", R, "money", 10),
    Column("Delivery", "delivery", R, "money", 10),
    Column("Stock", "stock", R, "text", 13),
    Column("Last checked", "checked", R, "datetime", 16),
    Column("Price check", "price_check", R, "text", 28),
    Column("Shop's barcode", "shop_barcode", R, "barcode", 16),
    Column("Hidden by", "hidden_by", R, "text", 11),
    Column("Show on site", "show", E, "yes_no", 12, "yes_no"),
    Column("Not this product", "not_this", A, "yes_no", 14, "yes_no"),
))

SHOP_ITEMS = Sheet("Shop items", "id", "link", (
    Column("Shop item id", "id", R, "id", 12),
    Column("Shop", "shop", R, "text", 20),
    Column("Shop's title", "title", R, "text", 50),
    Column("Link", "link", R, "link", 40),
    Column("Shop's barcode", "shop_barcode", R, "barcode", 16),
    Column("Price seen", "price", R, "money", 10),
    Column("Stock seen", "stock", R, "text", 13),
    Column("Last seen", "seen", R, "datetime", 16),
    Column("Found by", "found_by", R, "text", 14),
    Column("Status", "status", R, "text", 44),
    Column("Suggested id", "suggested", R, "int", 12),
    Column("Suggested product", "suggested_name", R, "text", 45),
    Column("Match %", "confidence", R, "int", 9),
    Column("Already shown there", "shown_there", R, "text", 40),
    Column("Link to product id", "link_to", A, "ref", 14),
    Column("Decision", "decision", A, "choice", 14, "decisions"),
))

# The key is the shop and the link together: these items have no row of their own on the site.
LEFT_OUT = Sheet("Left out", "link", "", (
    Column("Shop", "shop", R, "text", 20),
    Column("Shop's title", "title", R, "text", 50),
    Column("Link", "link", R, "link", 50),
    Column("Barcode given", "barcode", R, "barcode", 16),
    Column("Why left out", "why", R, "text", 44),
    Column("Link to product id", "link_to", A, "ref", 14),
    Column("Decision", "decision", A, "choice", 14, "left_decisions"),
))

SETS = Sheet("Sets", "id", "address", (
    Column("Set id", "id", R, "id", 10),
    Column("Address", "address", R, "link", 40),
    Column("Game", "game", R, "game", 18, "games"),   # read for new rows only
    Column("Name", "name", E, "text", 40),
    Column("Code", "code", E, "text", 10),
    Column("Release date", "release_date", E, "date", 13),
    Column("Date from", "date_from", R, "text", 14),
    Column("Products", "products", R, "int", 10),
))

SHEETS = (PRODUCTS, LISTINGS, SHOP_ITEMS, LEFT_OUT, SETS)
README = "Read me"
LISTS = "Lists"

YES_NO = ("Yes", "No")
DECISIONS = ("Link", "Not ours", "Look again")
LEFT_DECISIONS = ("Link", "Not ours")
# The hidden Lists sheet, one list a column, in this order.
LIST_NAMES = ("yes_no", "types", "games", "sets", "languages", "decisions", "left_decisions")

# The custom document property that carries the file number; Read me B2 carries it too.
REFERENCE_PROPERTY = "RipRaptor file"
REFERENCE = re.compile(r"RR-(\d+)-([0-9a-f]{12})")

README_LINES = (
    "Change only the orange columns. Grey columns come from the shops or the site and are ignored when you upload.",
    "Sort and filter as you like, but sort whole rows, never one column. Deleting a row changes nothing.",
    "To merge two pages of one product, put the id of the one to keep in Same as product id. "
    "For look-alikes that differ, put the other id in Not the same as.",
    "To add a product, add a row at the bottom of Products with new1, new2 and so on as its id, a game, a name, "
    "a type and a release date. Other sheets can use new1 as a product id.",
    "To link a shop's item, choose Link in Decision. Type a product id if the suggestion is wrong.",
    "To hide a product or one shop's price, set Show on site to No.",
    "To change a language, pick it in Language and the site adds it to the name.",
    "Barcodes stay as text. 5.05E+12 means Excel damaged a barcode.",
    "Prices, stock and shop titles come from the shops and are never changed from this file.",
    "Save as Excel Workbook (.xlsx) and upload it on the Your data page within 90 days. You see every change "
    "before anything happens, and you can undo the file.",
    "If Claude or someone else edits this file, ask them to change only orange columns and keep every id, then "
    "read the list of changes before you apply them.",
)

# Characters an Excel cell cannot hold (openpyxl refuses them); a shop's title sometimes carries one.
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
SPACES = re.compile(r"\s+")

STOCK = dict(Listing.Availability.choices)
PRICE_CHECK = dict(Listing.Sanity.choices)
FOUND_BY = dict(ShopProduct.Source.choices)
STATUS_ORDER = Case(
    When(status=ShopProduct.Status.REVIEW, then=Value(0)), When(status=ShopProduct.Status.LINKED, then=Value(1)),
    default=Value(2), output_field=IntegerField(),
)
WAITING = "Waiting for you"
WAITING_OTHER_PAGE = "Waiting for you: this shop's other page is already shown on the suggested product"
NOT_OURS = "Not ours"
LINKED_GONE = "Linked, not shown: its product was merged away, switched off or removed"
LINKED_NO_LISTING = "Linked, not shown: no listing"
LINKED_OTHER_PAGE = "Linked, not shown: this shop's other page is already shown on the product"
LINKED_WAITING = "Linked, waiting for the shop's next read"
TYPED_BY_YOU = "Type set by you"
TIDY_REMOVES = "No shop and no release date: the hourly tidy removes it"


def clean_text(value):
    """Text as the upload compares it: no control characters, no-break spaces as spaces, runs of spaces as
    one, no spaces at the ends. Applied alike to what is written and what comes back."""
    if value is None:
        return ""
    text = CONTROL.sub("", str(value)).replace("\xa0", " ")
    return SPACES.sub(" ", text).strip()


def shown_link(url, source_type):
    """A shop's link as the file shows it. eBay and Amazon links lose their query string, which carries
    tracking and our own tags, and says nothing about which product it is."""
    url = url or ""
    if source_type in MARKETPLACES:
        return url.split("?", 1)[0].split("#", 1)[0]
    return url


def digits(text):
    return "".join(ch for ch in text or "" if ch.isdigit())


def iso(day):
    return day.isoformat() if day else None


def local(moment):
    """A stored time as Excel holds one: London time, with no zone."""
    return timezone.localtime(moment).replace(tzinfo=None) if moment else None


def language_family(code):
    # "Chinese" that says no more is nearly always Simplified at UK shops, as duplicate_groups has it.
    return SIMPLIFIED if code == CHINESE else code


def lookalike_groups(products):
    """{product pk: group number} for products that may be one product under two pages, numbered in the order
    given; a product alone has none.

    ``products`` are (pk, game id, game slug, name, language, barcode) tuples, so this adds no query. Two
    products meet when they share the loose duplicate key within a game and language family, a Simplified
    Chinese set code within a game, or a barcode. The type is not part of any key: a "Booster Display" filed
    as a bundle is still the "Booster Box" it duplicates. This is where a wider key plugs in."""
    parent = {}

    def root(pk):
        while parent.get(pk, pk) != pk:
            parent[pk] = parent.get(parent[pk], parent[pk])
            pk = parent[pk]
        return pk

    first = {}
    order = []
    for pk, game_id, game_slug, name, language, ean in products:
        order.append(pk)
        keys = (
            ("loose", game_id, language_family(language), loose_key(name, game_slug)),
            ("code", game_id, chinese_code_key(name, game_slug)),
            ("ean", ean_key(ean)),
        )
        for key in keys:
            if not key[-1]:
                continue
            if key in first:
                a, b = root(first[key]), root(pk)
                if a != b:
                    parent[b] = a
            else:
                first[key] = pk
    members = defaultdict(int)
    for pk in order:
        members[root(pk)] += 1
    groups, numbers = {}, {}
    for pk in order:
        top = root(pk)
        if members[top] < 2:
            continue
        if top not in numbers:
            numbers[top] = len(numbers) + 1
        groups[pk] = numbers[top]
    return groups


def kept_apart_by_product():
    """{product pk: "455, 812"}, from the owner's standing Not the same answers: one query."""
    found = defaultdict(set)
    for pair in checks.kept_apart():
        a, b = sorted(pair)
        found[a].add(b)
        found[b].add(a)
    return {pk: ", ".join(str(other) for other in sorted(others)) for pk, others in found.items()}


def merged_into():
    """{slug of a merged-away product: the product kept}. merge_undoable repoints old aliases, so this is the
    final keeper."""
    return dict(ProductAlias.objects.values_list("slug", "product_id"))


def address_maker(name, *placeholders):
    """A fast reverse() for thousands of rows: the site address with placeholders to fill in."""
    path = reverse(name, args=[f"__{p}__" for p in placeholders])
    root = settings.RIPRAPTOR_SITE_URL

    def make(*values):
        url = path
        for placeholder, value in zip(placeholders, values):
            url = url.replace(f"__{placeholder}__", value)
        return root + url

    return make


# The rows ---------------------------------------------------------------------------------------------------
# Each generator yields (key, values, base): the row's key, the values written by column key, and the
# editable and integrity values as the upload is compared with them (type and language as codes, set as an
# id, barcode as digits, show on site as True or False, dates as ISO text).

def product_rows():
    """Every product except one merged away: switched off with its address now an old address of the product
    kept (the rule tidy_catalogue uses). By game, set (none last), type and name."""
    typed = types.owner_typed()
    apart = kept_apart_by_product()
    address = address_maker("web:product", "slug")
    rows = list(
        Product.objects.with_prices()
        .exclude(is_active=False, slug__in=ProductAlias.objects.values("slug"))
        .annotate(any_listings=Count("listings", distinct=True))
        .order_by("game__sort_order", "game__name", F("product_set__name").asc(nulls_last=True), "product_type",
                  "name", "pk")
        .values_list(
            "pk", "game_id", "game__name", "game__slug", "name", "product_type", "product_set_id",
            "product_set__name", "product_set__code", "language", "ean", "release_date", "is_active", "slug",
            "image", "created_at", "listing_count", "in_stock_count", "lowest_price", "any_listings",
        ).iterator(chunk_size=2000)
    )
    groups = lookalike_groups((row[0], row[1], row[3], row[4], row[9], row[10]) for row in rows)
    for (pk, _game_id, game, game_slug, name, kind, set_id, set_name, set_code, language, ean, released, active,
         slug, image, added, shops, in_stock, cheapest, any_listings) in rows:
        notes = []
        if pk in typed:
            notes.append(TYPED_BY_YOU)
        if not any_listings and released is None and not image:
            notes.append(TIDY_REMOVES)
        link = address(slug)
        values = {
            "id": pk, "address": link, "game": game, "name": name, "type": types.type_label(game_slug, kind),
            "set": set_name or "", "set_code": set_code or "", "language": languages.name(language),
            "barcode": ean or "", "release_date": released, "show": YES_NO[0] if active else YES_NO[1],
            "group": groups.get(pk), "apart": apart.get(pk, ""), "shops": shops, "in_stock": in_stock,
            "cheapest": cheapest, "added": local(added), "notes": "; ".join(notes),
        }
        base = {
            "address": link, "name": clean_text(name), "type": kind, "set": set_id, "language": language,
            "barcode": digits(ean), "release_date": iso(released), "show": bool(active),
        }
        yield pk, values, base


class ListingIndex:
    """What the Listings sheet saw, for the sheets after it: which product each shop's link is on, and the
    title and link of each shop's page on each product. Hidden listings count: they are still a shop's page."""

    def __init__(self):
        self.by_link = {}       # (retailer id, link key): product id
        self.by_product = {}    # (product id, retailer id): (title, link key)

    def add(self, retailer_id, url, product_id, title):
        key = link_key(url)
        self.by_link[(retailer_id, key)] = product_id
        self.by_product[(product_id, retailer_id)] = (title or url, key)


def listing_rows(index):
    """Every listing, hidden ones and ones on switched-off products included, by shop then product."""
    merged = merged_into()
    hidden_by = {}
    for listing_id, by_owner, ask_id in (
        CheckAnswer.objects.filter(kind=CheckAnswer.Kind.HIDE, undone_at__isnull=True, listing__isnull=False)
        .order_by("-created_at", "-pk").values_list("listing_id", "by_owner", "ask_id")
    ):
        hidden_by.setdefault(listing_id, "you" if by_owner else "Claude" if ask_id else "autopilot")
    rows = (
        Listing.objects.order_by("retailer__name", "product__name", "pk").values_list(
            "pk", "retailer_id", "retailer__name", "retailer__source_type", "title", "url", "product_id",
            "product__name", "product__is_active", "product__slug", "price", "delivery_cost", "delivery_known",
            "availability", "last_checked", "sanity", "sanity_reason", "shop_ean", "is_active",
        ).iterator(chunk_size=2000)
    )
    for (pk, retailer_id, shop, source, title, url, product_id, product, product_active, slug, price, delivery,
         delivery_known, stock, checked, sanity, reason, shop_ean, active) in rows:
        index.add(retailer_id, url, product_id, title)
        link = shown_link(url, source)
        if not product_active and slug in merged:
            shown = f"Merged into {merged[slug]}"
        else:
            shown = YES_NO[0] if product_active else YES_NO[1]
        if sanity == Listing.Sanity.EXCLUDED:
            check = f"Kept out: {reason}" if reason else "Kept out"
        else:
            check = PRICE_CHECK.get(sanity, sanity)
        values = {
            "id": pk, "shop": shop, "title": title, "link": link, "product": product_id, "product_name": product,
            "product_shown": shown, "price": price, "delivery": delivery if delivery_known else None,
            "stock": STOCK.get(stock, stock), "checked": local(checked), "price_check": check,
            "shop_barcode": shop_ean or "", "hidden_by": "" if active else hidden_by.get(pk, ""),
            "show": YES_NO[0] if active else YES_NO[1],
        }
        yield pk, values, {"link": link, "product": product_id, "show": bool(active)}


def shop_item_status(status, suggested_id, suggested_active, price, there, key):
    """What the Status column says, or None for a linked item the site shows, which is left out."""
    if status == ShopProduct.Status.REVIEW:
        return WAITING_OTHER_PAGE if there and there[1] != key else WAITING
    if status == ShopProduct.Status.IGNORED:
        return NOT_OURS
    if suggested_id is None or not suggested_active:
        return LINKED_GONE
    if there:
        return None if there[1] == key else LINKED_OTHER_PAGE
    return LINKED_WAITING if price is None else LINKED_NO_LISTING


def shop_item_rows(index):
    """Shop items waiting for a decision, ones he said are not ours, and linked ones the site does not show,
    by shop, then waiting, linked and not ours, then title. ``index`` comes from listing_rows."""
    rows = (
        ShopProduct.objects.alias(status_order=STATUS_ORDER)
        .order_by("retailer__name", "status_order", "title", "pk").values_list(
            "pk", "retailer_id", "retailer__name", "retailer__source_type", "title", "url", "shop_ean", "price",
            "availability", "last_seen", "source", "status", "suggested_id", "suggested__name",
            "suggested__is_active", "confidence",
        ).iterator(chunk_size=2000)
    )
    for (pk, retailer_id, shop, source_type, title, url, shop_ean, price, stock, seen, source, status, suggested_id,
         suggested, suggested_active, confidence) in rows:
        key = link_key(url)
        there = index.by_product.get((suggested_id, retailer_id)) if suggested_id else None
        text = shop_item_status(status, suggested_id, suggested_active, price, there, key)
        if text is None:
            continue
        link = shown_link(url, source_type)
        values = {
            "id": pk, "shop": shop, "title": title, "link": link, "shop_barcode": shop_ean or "", "price": price,
            "stock": STOCK.get(stock, stock or ""), "seen": local(seen), "found_by": FOUND_BY.get(source, source),
            "status": text, "suggested": suggested_id, "suggested_name": suggested or "", "confidence": confidence,
            "shown_there": there[0] if there else "",
        }
        yield pk, values, {"link": link, "status": status, "suggested": suggested_id}


def left_out_key(retailer_id, url):
    return f"{retailer_id}|{link_key(url)}"


def left_out_rows(retailer, index):
    """Items one shop's latest good whole read left out that look like sealed products of a game we list, less
    any the site already knows: a shop item row of that shop, or a listing at that shop with the same link.
    Only that one read's text is loaded."""
    text = (
        ImportRun.objects.filter(retailer=retailer, error="", note="", finished_at__isnull=False)
        .order_by("-started_at", "-pk").values_list("unmatched", flat=True).first()
    ) or ""
    known = {link_key(url) for url in ShopProduct.objects.filter(retailer=retailer).values_list("url", flat=True)}
    seen = set()
    for item in left_out_lines(text):
        key = link_key(item["url"])
        if key in known or key in seen or (retailer.pk, key) in index.by_link:
            continue
        seen.add(key)
        link = shown_link(item["url"], retailer.source_type)
        values = {"shop": retailer.name, "title": item["title"], "link": link, "barcode": item["barcode"],
                  "why": item["why"]}
        yield left_out_key(retailer.pk, item["url"]), values, {
            "link": link, "title": clean_text(item["title"]), "barcode": item["barcode"]}
        if len(seen) >= LEFT_OUT_LIMIT:
            break


def set_rows():
    """Every set, by game then newest release first."""
    address = address_maker("web:set", "game", "set")
    rows = (
        ProductSet.objects.annotate(products_count=Count("products", filter=Q(products__is_active=True)))
        .order_by("game__sort_order", "game__name", F("release_date").desc(nulls_last=True), "name", "pk")
        .values_list("pk", "game__name", "game__slug", "slug", "name", "code", "release_date",
                     "release_date_source", "products_count").iterator(chunk_size=2000)
    )
    for pk, game, game_slug, slug, name, code, released, source, products in rows:
        link = address(game_slug, slug)
        values = {
            "id": pk, "address": link, "game": game, "name": name, "code": code or "", "release_date": released,
            "date_from": "you" if source == OWNER else source, "products": products,
        }
        base = {"address": link, "name": clean_text(name), "code": clean_text(code), "release_date": iso(released)}
        yield pk, values, base


def list_values(set_names):
    """The hidden Lists sheet's columns, by LIST_NAMES."""
    labels = {label for _code, label in types.DEFAULT}
    for choices in types.GAME_TYPES.values():
        labels.update(label for _code, label in choices)
    return {
        "yes_no": list(YES_NO),
        "types": sorted(labels, key=str.casefold),
        "games": list(Game.objects.order_by("sort_order", "name").values_list("name", flat=True)),
        "sets": sorted(set(set_names), key=str.casefold),
        "languages": [languages.name(languages.ENGLISH), *languages.NAMES.values()],
        "decisions": list(DECISIONS),
        "left_decisions": list(LEFT_DECISIONS),
    }


# Writing ----------------------------------------------------------------------------------------------------

class Writer:
    """Cells as the sheets want them. Every text is a string cell, so a shop's title that starts with = is never
    a formula; editable cells are text-formatted, so a barcode or an id he types stays as typed."""

    def __init__(self):
        from openpyxl.styles import Font, PatternFill

        self.bold = Font(bold=True)
        # Opaque: openpyxl would otherwise write the colour with a zero alpha.
        self.editable = PatternFill("solid", start_color="FF" + EDITABLE_FILL, end_color="FF" + EDITABLE_FILL)
        self.grey = PatternFill("solid", start_color="FF" + READ_ONLY_FILL, end_color="FF" + READ_ONLY_FILL)

    def cell(self, ws, value, number_format=None, link=False):
        from openpyxl.cell import WriteOnlyCell

        if isinstance(value, str):
            value = CONTROL.sub("", value)
            if not value:
                return None
        cell = WriteOnlyCell(ws, value=value)
        if isinstance(value, str):
            cell.data_type = "s"
            if link and value.startswith(("http://", "https://")):
                cell.hyperlink = value
        if number_format:
            cell.number_format = number_format
        return cell

    def value(self, ws, column, value):
        if value is None or value == "":
            return None
        kind = column.kind
        if kind == "date":
            return self.cell(ws, value, DATE_FORMAT)
        if kind == "datetime":
            return self.cell(ws, value, DATETIME_FORMAT)
        if kind == "money":
            return self.cell(ws, value, MONEY_FORMAT)
        text_format = TEXT_FORMAT if column.mark != READ_ONLY else None
        if isinstance(value, str):
            return self.cell(ws, value, text_format, link=kind == "link")
        if text_format:
            return self.cell(ws, value, text_format)
        return value

    def heading(self, ws, column):
        cell = self.cell(ws, column.heading)
        cell.font = self.bold
        if column.mark != READ_ONLY:
            cell.fill = self.editable
        return cell

    def prepare(self, ws, sheet):
        """What Excel must see before the first row: widths, frozen headings, text columns."""
        from openpyxl.utils import get_column_letter

        ws.freeze_panes = "B2"
        for number, column in enumerate(sheet.columns, start=1):
            dimension = ws.column_dimensions[get_column_letter(number)]
            dimension.width = column.width
            if column.mark != READ_ONLY:
                dimension.number_format = DATE_FORMAT if column.kind == "date" else TEXT_FORMAT
        ws.append([self.heading(ws, column) for column in sheet.columns])

    def finish(self, ws, sheet, rows, lists):
        """After the last row: grey read-only columns, dropdowns and the filter."""
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation

        last = rows + 1
        for number, column in enumerate(sheet.columns, start=1):
            letter = get_column_letter(number)
            if column.mark == READ_ONLY:
                # One rule a column: a fill on each of 13,000 cells took 8.7 seconds.
                ws.conditional_formatting.add(f"{letter}1:{letter}{last}", FormulaRule(formula=["TRUE"], fill=self.grey))
            if column.dropdown:
                at = LIST_NAMES.index(column.dropdown)
                list_letter = get_column_letter(at + 1)
                size = max(len(lists[column.dropdown]), 1)
                # A warning, not a stop: a set added on the Sets sheet in the same file is not in the list yet.
                check = DataValidation(
                    type="list", formula1=f"{LISTS}!${list_letter}$1:${list_letter}${size}", allow_blank=True,
                    errorStyle="warning", showErrorMessage=True, errorTitle="Not in the list",
                    error="This is not one of the choices. Keep it only if you meant it.",
                )
                check.add(f"{letter}2:{letter}{last + SPARE_ROWS}")
                ws.data_validations.append(check)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(sheet.columns))}{last}"


def build_export(export):
    """(workbook bytes, base, rows per sheet) for one DataExport, which must be saved: its number goes in the
    file. Reads each table once, row by row, outside any transaction: a read inside atomic() would take
    SQLite's write lock and stall every visitor while the file is made."""
    from openpyxl import Workbook
    from openpyxl.packaging.custom import StringProperty

    made = timezone.now()
    wb = Workbook(write_only=True)
    wb.properties.creator = "RipRaptor"
    wb.properties.title = "RipRaptor data"
    readme = wb.create_sheet(README)
    writer = Writer()
    retailer = export.left_out_retailer
    sheets = [PRODUCTS, LISTINGS, SHOP_ITEMS, *([LEFT_OUT] if retailer else []), SETS]
    pages = {sheet.name: wb.create_sheet(sheet.name) for sheet in sheets}
    lists_ws = wb.create_sheet(LISTS)
    lists_ws.sheet_state = "hidden"

    index = ListingIndex()
    set_names = []

    def sets():
        for key, values, base in set_rows():
            set_names.append(values["name"])
            yield key, values, base

    makers = {
        PRODUCTS.name: product_rows,
        LISTINGS.name: lambda: listing_rows(index),
        SHOP_ITEMS.name: lambda: shop_item_rows(index),
        LEFT_OUT.name: lambda: left_out_rows(retailer, index),
        SETS.name: sets,
    }
    base = {"v": 1, "sheets": {}}
    counts, written = {}, {}
    for sheet in sheets:
        ws = pages[sheet.name]
        writer.prepare(ws, sheet)
        rows = base["sheets"][sheet.name] = {}
        for key, values, row_base in makers[sheet.name]():
            ws.append([writer.value(ws, column, values.get(column.key)) for column in sheet.columns])
            rows[str(key)] = row_base
        written[sheet.name] = counts[sheet.name] = len(rows)

    # Product set names come from the Sets sheet, so the lists and dropdowns follow the rows.
    lists = list_values(set_names)
    for sheet in sheets:
        writer.finish(pages[sheet.name], sheet, written[sheet.name], lists)
    longest = max(len(values) for values in lists.values())
    for row in range(longest):
        lists_ws.append([writer.cell(lists_ws, lists[name][row]) if row < len(lists[name]) else None
                         for name in LIST_NAMES])

    reference = export.reference
    readme.column_dimensions["A"].width = 16
    readme.column_dimensions["B"].width = 60
    title = writer.cell(readme, "RipRaptor data")
    title.font = writer.bold
    readme.append([title])
    readme.append([writer.cell(readme, "File reference"), writer.cell(readme, reference)])
    readme.append([writer.cell(readme, "Downloaded"), writer.cell(readme, local(made), DATETIME_FORMAT)])
    readme.append([writer.cell(readme, "Rows per sheet"),
                   writer.cell(readme, ", ".join(f"{name} {count:,}" for name, count in counts.items()))])
    readme.append([])
    for line in README_LINES:
        readme.append([writer.cell(readme, line)])
    wb.custom_doc_props.append(StringProperty(name=REFERENCE_PROPERTY, value=reference))

    with SpooledTemporaryFile(max_size=SPOOL_BYTES) as out:
        wb.save(out)
        out.seek(0)
        data = out.read()
    return data, base, counts


def quick_reference(archive):
    """(export pk, token) from an uploaded workbook's custom properties, or None, without openpyxl: the upload
    page refuses a file with no number, or an old one, before anything else is done. ``archive`` is an open
    zipfile.ZipFile."""
    try:
        text = archive.read("docProps/custom.xml").decode("utf-8", "replace")
    except KeyError:
        return None
    found = REFERENCE.search(text)
    return (int(found.group(1)), found.group(2)) if found else None
